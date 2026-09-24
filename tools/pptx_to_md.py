"""把作图方案 PPT 拆成大模型好读的格式：一份 Markdown + 图片。

输出目录结构：
    <out>/brief.md          每页的文字、表格、页内图片清单
    <out>/slides/           每页整页截图（需要装 LibreOffice）
    <out>/images/           每页里嵌入的原图（参考图、竞品图等）

用法：
    python tools/pptx_to_md.py 方案.pptx -o briefs/xxx --title "女装-裤子-弯刀蕾丝阔腿裤"
"""

import argparse
import hashlib
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

SOFFICE_CANDIDATES = [
    "soffice",
    "libreoffice",
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
]


def find_soffice(explicit=None):
    for cand in [explicit] if explicit else SOFFICE_CANDIDATES:
        path = shutil.which(cand) or (cand if cand and Path(cand).exists() else None)
        if path:
            return path
    return None


def render_slides(pptx_path, out_dir, soffice, width_px):
    """PPT -> PDF (LibreOffice) -> 每页一张 JPG (PyMuPDF)。返回生成的文件名列表。"""
    import pymupdf as fitz

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        src = tmp / "deck.pptx"
        shutil.copyfile(pptx_path, src)
        profile = (tmp / "lo_profile").as_uri()
        subprocess.run(
            [soffice, f"-env:UserInstallation={profile}", "--headless",
             "--convert-to", "pdf", "--outdir", str(tmp), str(src)],
            check=True, capture_output=True, timeout=600,
        )
        pdf = tmp / "deck.pdf"
        if not pdf.exists():
            raise RuntimeError("LibreOffice 没有生成 PDF，请确认装了 Impress 组件")

        out_dir.mkdir(parents=True, exist_ok=True)
        names = []
        with fitz.open(pdf) as doc:
            for i, page in enumerate(doc, 1):
                zoom = width_px / page.rect.width
                pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
                name = f"slide-{i:02d}.jpg"
                pix.save(out_dir / name, jpg_quality=85)
                names.append(name)
        return names


def iter_shapes(shapes, parent_box=None):
    """展开组合形状。组合里的子形状用组合整体的位置，方便判断在页面哪一侧。"""
    for sh in shapes:
        box = parent_box or (sh.left or 0, sh.top or 0, sh.width or 0, sh.height or 0)
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(sh.shapes, box)
        else:
            yield sh, box


def reading_order(items, slide_h):
    # 先按从上到下分行（每行约占页面高度的 6%），行内从左到右
    band = max(slide_h * 0.06, 1)
    return sorted(items, key=lambda it: (int(it[1][1] // band), it[1][0]))


def describe_position(box, slide_w, slide_h):
    left, top, width, height = box
    cx = (left + width / 2) / slide_w
    cy = (top + height / 2) / slide_h
    horiz = "左侧" if cx < 0.36 else "右侧" if cx > 0.64 else "中间"
    vert = "上部" if cy < 0.36 else "下部" if cy > 0.64 else "中部"
    return f"{horiz}{vert}", f"宽 {width / slide_w:.0%} × 高 {height / slide_h:.0%}"


def text_of(shape):
    lines = []
    for para in shape.text_frame.paragraphs:
        parts = []
        for run in para.runs:
            txt = run.text.replace("\x0b", "\n")
            url = run.hyperlink.address if run.hyperlink else None
            parts.append(f"[{txt}]({url})" if url and txt.strip() and txt.strip() != url else txt)
        lines.append("".join(parts).rstrip())
    # 去掉首尾空行，中间连续空行合并成一个
    out, blank = [], False
    for line in lines:
        if line.strip():
            out.append(line)
            blank = False
        elif out and not blank:
            out.append("")
            blank = True
    return "\n".join(out).strip()


def table_md(table):
    rows = [[c.text.replace("\n", " ").replace("|", "\\|").strip() for c in r.cells] for r in table.rows]
    if not rows:
        return ""
    md = ["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * len(rows[0])]
    md += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(md)


def crop_note(shape):
    try:
        crops = [shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom]
    except (AttributeError, ValueError):
        return ""
    return "（页面上只显示了原图的一部分）" if any(abs(c) > 0.01 for c in crops) else ""


def convert(pptx_path, out_dir, title=None, render=True, soffice=None, width_px=1600):
    pptx_path, out_dir = Path(pptx_path), Path(out_dir)
    prs = Presentation(pptx_path)
    sw, sh = prs.slide_width, prs.slide_height
    img_dir = out_dir / "images"
    img_dir.mkdir(parents=True, exist_ok=True)

    slide_imgs = []
    if render:
        exe = find_soffice(soffice)
        if exe:
            slide_imgs = render_slides(pptx_path, out_dir / "slides", exe, width_px)
        else:
            print("提示：没找到 LibreOffice，跳过整页截图（只导出文字和原图）", file=sys.stderr)

    title = title or pptx_path.stem
    md = [
        f"# {title}",
        "",
        f"- 来源文件：`{pptx_path.name}`",
        f"- 共 {len(prs.slides)} 页",
        "- 阅读方法：每页先看「整页截图」了解排版和标注，再看「文字」和「页内图片」。"
        "页内图片是 PPT 里放的原图（多为竞品参考图），位置描述指它在这一页的哪个区域。",
        "",
    ]

    for idx, slide in enumerate(prs.slides, 1):
        md += ["---", "", f"## 第 {idx} 页", ""]
        if idx <= len(slide_imgs):
            md += [f"![第 {idx} 页整页截图](slides/{slide_imgs[idx - 1]})", ""]

        texts, tables, pics = [], [], []
        for shape, box in iter_shapes(slide.shapes):
            if shape.has_text_frame:
                t = text_of(shape)
                if t:
                    is_title = shape.is_placeholder and ("标题" in shape.name or "title" in shape.name.lower())
                    texts.append((t, box, is_title))
            if getattr(shape, "has_table", False) and shape.has_table:
                tables.append((table_md(shape.table), box))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):
                try:
                    image = shape.image
                except (AttributeError, ValueError):
                    continue
                pics.append(((shape, image), box))

        if texts:
            md += ["### 文字", ""]
            for t, _, is_title in reading_order(texts, sh):
                md += [f"**{t}**" if is_title else t, ""]

        if tables:
            md += ["### 表格", ""]
            for t, _ in reading_order(tables, sh):
                md += [t, ""]

        if pics:
            md += ["### 页内图片", "", "| 图片 | 在页面的位置 | 占页面大小 | 原图尺寸 |", "| --- | --- | --- | --- |"]
            for n, ((shape, image), box) in enumerate(reading_order(pics, sh), 1):
                name = f"s{idx:02d}-{n}.{image.ext}"
                (img_dir / name).write_bytes(image.blob)
                pos, size = describe_position(box, sw, sh)
                px = "×".join(map(str, image.size)) if image.size else "未知"
                md.append(f"| ![{name}](images/{name}) | {pos}{crop_note(shape)} | {size} | {px} |")
            md.append("")

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                md += ["### 备注", "", notes, ""]

    (out_dir / "brief.md").write_text("\n".join(md), encoding="utf-8")
    return out_dir / "brief.md"


def main():
    ap = argparse.ArgumentParser(description="把 PPT 拆成 Markdown + 图片，给大模型读")
    ap.add_argument("pptx")
    ap.add_argument("-o", "--out", help="输出目录，默认和 PPT 同名")
    ap.add_argument("--title", help="文档标题，默认用 PPT 文件名")
    ap.add_argument("--no-render", action="store_true", help="不生成整页截图")
    ap.add_argument("--soffice", help="LibreOffice 程序路径（自动找不到时填写）")
    ap.add_argument("--width", type=int, default=1600, help="整页截图宽度，默认 1600 像素")
    args = ap.parse_args()

    out = args.out or os.path.splitext(args.pptx)[0]
    path = convert(args.pptx, out, args.title, not args.no_render, args.soffice, args.width)
    print(f"完成：{path}")


if __name__ == "__main__":
    main()
